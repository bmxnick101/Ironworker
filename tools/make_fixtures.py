"""Build tests/fixtures.json: copy Nick's exact O3/P3/Q3 formulas into many rows,
recalculate with LibreOffice, read back the results."""
import openpyxl, json, subprocess, shutil, os, itertools
src = openpyxl.load_workbook("tools/source.xlsx").active
fO, fP, fQ = src["O3"].value, src["P3"].value, src["Q3"].value
wb = openpyxl.Workbook(); ws = wb.active
mats = ["Steel", "Stainless Steel", "Aluminum"]
thks = [0.0598, 0.075, 0.1, 0.109, 0.125, 0.1875, 0.25, 0.3125, 0.375, 0.5, 0.625, 0.75, 1.0]
punches = [0.1875, 0.3125, 0.4375, 0.5, 0.6875, 0.8125, 1.0, 1.125]
r = 3
for m, t, p in itertools.product(mats, thks, punches):
    ws[f"L{r}"], ws[f"M{r}"], ws[f"N{r}"] = m, t, p
    for col, f in (("O", fO), ("P", fP), ("Q", fQ)):
        ws[f"{col}{r}"] = f.replace("L3", f"L{r}").replace("M3", f"M{r}").replace("N3", f"N{r}").replace("O3", f"O{r}").replace("mround", "MROUND").replace("IFS(", "_xlfn.IFS(")
    r += 1
os.makedirs("/tmp/fx", exist_ok=True)
wb.save("/tmp/fx/in.xlsx")
subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", "/tmp/fx/out", "/tmp/fx/in.xlsx"], check=True, capture_output=True)
res = openpyxl.load_workbook("/tmp/fx/out/in.xlsx", data_only=True).active
rows = [{"material": res[f"L{i}"].value, "thickness": res[f"M{i}"].value, "punch": res[f"N{i}"].value,
         "multiplier": res[f"O{i}"].value, "die": res[f"P{i}"].value, "tons": res[f"Q{i}"].value} for i in range(3, r)]
orig = openpyxl.load_workbook("tools/source.xlsx", data_only=True).active
cached = {"material": orig["L3"].value, "thickness": orig["M3"].value, "punch": orig["N3"].value,
          "multiplier": orig["O3"].value, "die": orig["P3"].value, "tons": orig["Q3"].value}
json.dump({"cachedFromNicksFile": cached, "recalculatedWithHisFormulas": rows}, open("tests/fixtures.json", "w"), indent=0)
print(len(rows), "rows; sample", rows[0], "cached", cached)
assert all(isinstance(x["tons"], float) and x["die"] is not None for x in rows)
