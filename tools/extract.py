"""Regenerate data/tooling.json + data/sizes.json from Nick's spreadsheet (Sheet1).
Usage: python3 tools/extract.py tools/source.xlsx"""
import json, re, sys
from fractions import Fraction
import openpyxl
src = sys.argv[1] if len(sys.argv) > 1 else "tools/source.xlsx"
ws = openpyxl.load_workbook(src).active
def frac(label):
    s = label.replace('"', '').strip()
    whole, _, rest = s.partition('-') if '-' in s else ('0', '', s)
    if '/' not in s: return float(s)
    return float(int(whole) + Fraction(rest))
# Confirmed by Nick (2026-10-02), matching the sheet headers and the Cleveland Punch & Die catalog:
# F16 and F26 are PUNCHES; F55 and F62 are DIES.
GROUP_KIND = {"F16": "punch", "F26": "punch", "F55": "die", "F62": "die"}
groups = []
for col in "BDFH":
    h = ws[f"{col}1"].value
    m = re.match(r"(.*?)\s+(F\d+)$", h.strip())
    code = m.group(2)
    kind = GROUP_KIND[code]
    groups.append({"column": col, "header": h, "code": code, "kind": kind,
                   "name": f"{code} " + ("Punches" if kind == "punch" else "Dies"), "headerSays": "punch" if "PUNCH" in h.upper() else "die"})
sizes, items = [], []
for r in range(2, ws.max_row + 1):
    lab = ws[f"A{r}"].value
    if not lab: continue
    d = frac(lab); sizes.append({"label": lab, "in": d, "row": r})
    for g in groups:
        v = ws[f"{g['column']}{r}"].value
        if v:
            q = re.search(r"\((\d+)\)", str(v))
            items.append({"id": f"{g['code']}-{r}", "kind": g["kind"], "group": g["code"], "shape": "round",
                          "size": d, "sizeLabel": lab, "qty": int(q.group(1)) if q else None,
                          "location": "", "notes": "", "inStock": True,
                          "source": "sheet", "sourceCell": f"{g['column']}{r}", "mark": str(v)})
json.dump({"source": "Sheet1 (columns A,B,D,F,H)", "note": "Only cells marked X are listed; an 'X' means available. Qty only where the sheet gives one, e.g. 'X (2)'. Location is not in the sheet. Groups: F16 and F26 are punches, F55 and F62 are dies (matches the sheet headers).",
           "groupKind": GROUP_KIND,
           "groups": groups, "items": items}, open("data/tooling.json", "w"), indent=1)
json.dump({"source": "Sheet1 column A", "sizes": sizes}, open("data/sizes.json", "w"), indent=1)
print(len(sizes), "sizes;", len(items), "tooling items")
